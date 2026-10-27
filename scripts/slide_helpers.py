"""
AgroVision-Mine: Complete Final Project Review Presentation Builder
Generates AgroVision-Mine_Complete_Final_Project_Review.pptx
"""

import os
import sys
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --- COLOR PALETTE (From Review 1 Template) ---
DARK_GREEN = RGBColor(27, 67, 50)      # #1B4332 Primary Brand
FOREST_GREEN = RGBColor(45, 106, 79)    # #2D6A4F Medium Forest
MINT_ACCENT = RGBColor(116, 198, 157)   # #74C69D Accent Green
MINT_TINT = RGBColor(216, 243, 220)     # #D8F3DC Light Green Card
CARD_BG = RGBColor(248, 249, 250)       # #F8F9FA Off-white / Gray Card
BORDER_GRAY = RGBColor(220, 225, 230)   # #DCE1E6 Subtle Border
TEXT_MAIN = RGBColor(26, 26, 26)        # #1A1A1A Main Body Text
TEXT_MUTED = RGBColor(90, 100, 110)     # #5A646E Muted Text
TEXT_WHITE = RGBColor(255, 255, 255)    # #FFFFFF White
AMBER_WARN = RGBColor(180, 83, 9)       # Warning Amber
NAVY_BLUE = RGBColor(30, 58, 138)       # Deep Navy Accent

FONT_TITLE = "Bookman Old Style"
FONT_BODY = "Bookman Old Style"
FONT_CODE = "Consolas"

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)


def create_blank_slide(prs):
    blank_layout = prs.slide_layouts[6] # blank
    slide = prs.slides.add_slide(blank_layout)
    return slide


def add_slide_header(slide, title_text, category_tag="CSI3010: DATA WAREHOUSING & DATA MINING"):
    """Adds a standard academic header with category tag and title."""
    # Category Tag
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = category_tag.upper()
    p_tag.font.name = FONT_BODY
    p_tag.font.size = Pt(11)
    p_tag.font.bold = True
    p_tag.font.color.rgb = FOREST_GREEN

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.65))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = FONT_TITLE
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = DARK_GREEN


def add_slide_footer(slide, slide_num, total_slides=24):
    """Adds a subtle standard footer bar."""
    # Divider line
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.8), Inches(7.0), Inches(11.733), Inches(0.02)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = BORDER_GRAY
    line.line.color.rgb = BORDER_GRAY

    # Footer text left
    foot_left = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(9.0), Inches(0.3))
    tf_l = foot_left.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = tf_l.margin_top = tf_l.margin_right = tf_l.margin_bottom = 0
    p_l = tf_l.paragraphs[0]
    p_l.text = "AgroVision-Mine: Multimodal Data Warehouse & Visual Pattern Mining Framework"
    p_l.font.name = FONT_BODY
    p_l.font.size = Pt(9.5)
    p_l.font.color.rgb = TEXT_MUTED

    # Footer text right (Slide Number)
    foot_right = slide.shapes.add_textbox(Inches(10.5), Inches(7.05), Inches(2.0), Inches(0.3))
    tf_r = foot_right.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = tf_r.margin_top = tf_r.margin_right = tf_r.margin_bottom = 0
    p_r = tf_r.paragraphs[0]
    p_r.alignment = PP_ALIGN.RIGHT
    p_r.text = f"{slide_num} / {total_slides}"
    p_r.font.name = FONT_BODY
    p_r.font.size = Pt(9.5)
    p_r.font.bold = True
    p_r.font.color.rgb = FOREST_GREEN


def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=BORDER_GRAY):
    """Creates a card shape with rounded rectangle and background."""
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        left, top, width, height
    )
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
    else:
        card.line.fill.background()
    return card


print("Helper module loaded successfully.")
