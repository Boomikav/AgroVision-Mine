import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = Presentation("AgroVision-Mine_Final_Review_AUDITED.pptx")

# 1. Slide 11: Polish TextBox 9 formatting
s11 = prs.slides[10]
for shape in s11.shapes:
    if shape.name == "TextBox 9" and shape.has_text_frame:
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.1)
        tf.margin_bottom = Inches(0.1)
        tf.margin_left = Inches(0.1)
        tf.margin_right = Inches(0.1)
        
        # Format title
        p0 = tf.paragraphs[0]
        p0.font.size = Pt(13)
        p0.font.bold = True
        p0.font.color.rgb = RGBColor(23, 93, 68)
        p0.space_after = Pt(8)
        
        # Format bullets
        for p in tf.paragraphs[1:]:
            p.font.size = Pt(10.5)
            p.font.color.rgb = RGBColor(15, 31, 24)
            p.space_after = Pt(6)
            p.space_before = Pt(0)

# 2. Slide 15: Polish Table 10 and Note layout
s15 = prs.slides[14]
for shape in s15.shapes:
    if shape.has_table:
        table = shape.table
        shape.top = 1850000
        shape.height = 2300000
        for r_idx, row in enumerate(table.rows):
            row.height = Pt(24)
            for cell in row.cells:
                cell.margin_top = Inches(0.04)
                cell.margin_bottom = Inches(0.04)
                cell.margin_left = Inches(0.06)
                cell.margin_right = Inches(0.06)
                for p in cell.text_frame.paragraphs:
                    p.font.size = Pt(8.5) if r_idx > 0 else Pt(9)
                    if r_idx == 0:
                        p.font.bold = True
    elif shape.name in ["Rounded Rectangle 11", "TextBox 12"]:
        if shape.name == "Rounded Rectangle 11":
            shape.top = 4350000
            shape.height = 1500000
        elif shape.name == "TextBox 12" and shape.has_text_frame:
            shape.top = 4400000
            shape.height = 1400000
            for p in shape.text_frame.paragraphs:
                p.font.size = Pt(8.5)
    elif shape.shape_type == 13: # Picture 13
        shape.top = 1850000
        shape.height = 2400000

prs.save("AgroVision-Mine_Final_Review_AUDITED.pptx")
print("Saved polished AgroVision-Mine_Final_Review_AUDITED.pptx!")
