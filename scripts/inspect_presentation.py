import sys
from pptx import Presentation

pptx_path = "AgroVision-Mine_Complete_Final_Project_Review(1).pptx"
prs = Presentation(pptx_path)

print(f"Total slides: {len(prs.slides)}")

for idx, slide in enumerate(prs.slides):
    title = "NO TITLE"
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text:
            text = shape.text_frame.text.strip().replace('\n', ' ')
            if "slide" in text.lower() or idx == 0 or len(text) < 80:
                title = text[:60]
                break
    print(f"Slide {idx + 1}: {title} (shapes: {len(slide.shapes)})")
