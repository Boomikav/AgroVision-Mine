import os
from pptx import Presentation

paths = [
    r"AgroVision-Mine_Complete_Final_Project_Review.pptx",
    r"C:\Users\ajaya\Downloads\AgroVision-Mine_Complete_Final_Project_Review_REVISED (1).pptx",
]

for p in paths:
    if os.path.exists(p):
        prs = Presentation(p)
        print(f"\n--- {p} ---")
        print(f"Slide count: {len(prs.slides)}")
        for i, s in enumerate(prs.slides):
            texts = [shp.text_frame.text.strip().replace('\n', ' ') for shp in s.shapes if shp.has_text_frame and shp.text_frame.text]
            first_text = texts[0][:50] if texts else "EMPTY"
            second_text = texts[1][:50] if len(texts) > 1 else ""
            print(f"  Slide {i+1}: {first_text} | {second_text}")
