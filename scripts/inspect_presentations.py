"""Deep inspection of TermiteSense-IoT and AgroVision-Mine presentations."""
import sys
from pathlib import Path
import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt

def inspect_slide(slide, idx):
    data = {"index": idx + 1, "shapes": []}
    for shape in slide.shapes:
        s_info = {
            "name": shape.name,
            "type": str(shape.shape_type),
            "top": round(shape.top / 914400, 2),
            "left": round(shape.left / 914400, 2),
            "width": round(shape.width / 914400, 2),
            "height": round(shape.height / 914400, 2),
        }
        if shape.has_text_frame:
            tf = shape.text_frame
            s_info["text"] = tf.text.strip()
            paras = []
            for p in tf.paragraphs:
                p_text = p.text.strip()
                if p_text:
                    fn = p.font.name or (p.runs[0].font.name if p.runs and p.runs[0].font.name else None)
                    fs = p.font.size.pt if p.font.size else (p.runs[0].font.size.pt if p.runs and p.runs[0].font.size else None)
                    color = None
                    try:
                        if p.font.color and p.font.color.rgb:
                            color = str(p.font.color.rgb)
                    except:
                        pass
                    if not color and p.runs:
                        try:
                            if p.runs[0].font.color and p.runs[0].font.color.rgb:
                                color = str(p.runs[0].font.color.rgb)
                        except:
                            pass
                    bold = p.font.bold if p.font.bold is not None else (p.runs[0].font.bold if p.runs else None)
                    paras.append({"text": p_text[:100], "font": fn, "size": fs, "color": color, "bold": bold})
            s_info["paragraphs"] = paras
        elif shape.has_table:
            t = shape.table
            rows_data = []
            for r in t.rows:
                rows_data.append([c.text.strip().replace("\n", " ") for c in r.cells])
            s_info["table"] = {"rows": len(t.rows), "cols": len(t.columns), "data": rows_data}
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            s_info["picture"] = True
        
        # Check fill color
        try:
            if shape.fill.type:
                s_info["fill"] = str(shape.fill.fore_color.rgb) if shape.fill.fore_color else None
        except:
            pass
            
        data["shapes"].append(s_info)
    return data

def run_inspection():
    out_dir = Path("data/reports/presentation_inspection")
    out_dir.mkdir(parents=True, exist_ok=True)
    
    # 1. TermiteSense
    ts_prs = pptx.Presentation("TermiteSense-IoT_Faculty_Review_Refined.pptx")
    with open(out_dir / "termitesense_inspection.txt", "w", encoding="utf-8") as f:
        f.write(f"TERMITESENSE TOTAL SLIDES: {len(ts_prs.slides)}\n")
        f.write(f"SLIDE DIMS: {ts_prs.slide_width/914400:.2f} x {ts_prs.slide_height/914400:.2f}\n\n")
        for i, s in enumerate(ts_prs.slides):
            data = inspect_slide(s, i)
            f.write(f"=== SLIDE {data['index']} ===\n")
            for shape in data["shapes"]:
                if "text" in shape and shape["text"]:
                    f.write(f"  Shape '{shape['name']}' ({shape['left']}\", {shape['top']}\", {shape['width']}\"x{shape['height']}\"):\n")
                    for p in shape["paragraphs"]:
                        f.write(f"    - [{p['font']}|{p['size']}pt|col:{p['color']}|b:{p['bold']}] {p['text']}\n")
                elif "table" in shape:
                    f.write(f"  Table '{shape['name']}' ({shape['table']['rows']}x{shape['table']['cols']}):\n")
                    for row in shape["table"]["data"][:4]:
                        f.write(f"    {row}\n")
                elif shape.get("picture"):
                    f.write(f"  Picture '{shape['name']}' ({shape['left']}\", {shape['top']}\", {shape['width']}\"x{shape['height']}\")\n")
            f.write("\n")
            
    # 2. AgroVision-Mine
    av_prs = pptx.Presentation("AgroVision-Mine_Final_Review_AUDITED(2).pptx")
    with open(out_dir / "agrovision_inspection.txt", "w", encoding="utf-8") as f:
        f.write(f"AGROVISION TOTAL SLIDES: {len(av_prs.slides)}\n")
        f.write(f"SLIDE DIMS: {av_prs.slide_width/914400:.2f} x {av_prs.slide_height/914400:.2f}\n\n")
        for i, s in enumerate(av_prs.slides):
            data = inspect_slide(s, i)
            f.write(f"=== SLIDE {data['index']} ===\n")
            for shape in data["shapes"]:
                if "text" in shape and shape["text"]:
                    f.write(f"  Shape '{shape['name']}' ({shape['left']}\", {shape['top']}\", {shape['width']}\"x{shape['height']}\"):\n")
                    for p in shape["paragraphs"]:
                        f.write(f"    - [{p['font']}|{p['size']}pt|col:{p['color']}|b:{p['bold']}] {p['text']}\n")
                elif "table" in shape:
                    f.write(f"  Table '{shape['name']}' ({shape['table']['rows']}x{shape['table']['cols']}):\n")
                    for row in shape["table"]["data"]:
                        f.write(f"    {row}\n")
                elif shape.get("picture"):
                    f.write(f"  Picture '{shape['name']}' ({shape['left']}\", {shape['top']}\", {shape['width']}\"x{shape['height']}\")\n")
            f.write("\n")

    print("Inspection complete. Files saved to data/reports/presentation_inspection/")

if __name__ == "__main__":
    run_inspection()
