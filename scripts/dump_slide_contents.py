import os
from pptx import Presentation

prs = Presentation("AgroVision-Mine_Complete_Final_Project_Review.pptx")

with open("docs/presentation_content_dump.txt", "w", encoding="utf-8") as out:
    for idx, slide in enumerate(prs.slides):
        out.write(f"\n{'='*70}\n")
        out.write(f"SLIDE {idx + 1} (Shapes: {len(slide.shapes)})\n")
        out.write(f"{'='*70}\n")
        
        for s_idx, shape in enumerate(slide.shapes):
            name = shape.name
            shape_type = str(shape.shape_type)
            has_text = shape.has_text_frame
            has_table = shape.has_table
            
            out.write(f"  Shape {s_idx + 1}: Name='{name}', Type={shape_type}\n")
            if has_text:
                text = shape.text_frame.text.strip()
                if text:
                    for p in shape.text_frame.paragraphs:
                        if p.text.strip():
                            out.write(f"    [P]: {p.text.strip()}\n")
            elif has_table:
                table = shape.table
                out.write(f"    [TABLE {len(table.rows)}x{len(table.columns)}]:\n")
                for r_idx, row in enumerate(table.rows):
                    row_vals = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                    out.write(f"      R{r_idx+1}: {' | '.join(row_vals)}\n")
            elif shape.shape_type == 13: # Picture
                out.write(f"    [PICTURE]: left={shape.left}, top={shape.top}, width={shape.width}, height={shape.height}\n")

print("Dumped presentation contents to docs/presentation_content_dump.txt")
