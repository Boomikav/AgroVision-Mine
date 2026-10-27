import docx
from docx.oxml.ns import qn

doc = docx.Document('Paper_template-a4 (1).docx')
print('=== SECTIONS ===')
for idx, sec in enumerate(doc.sections):
    print(f'Section {idx}: start_type={sec.start_type}')
    print(f'  page_width={sec.page_width.pt} pt ({sec.page_width.mm:.1f} mm), page_height={sec.page_height.pt} pt ({sec.page_height.mm:.1f} mm)')
    print(f'  margins: top={sec.top_margin.pt} pt ({sec.top_margin.mm:.1f} mm), bottom={sec.bottom_margin.pt} pt ({sec.bottom_margin.mm:.1f} mm), left={sec.left_margin.pt} pt ({sec.left_margin.mm:.1f} mm), right={sec.right_margin.pt} pt ({sec.right_margin.mm:.1f} mm)')
    sectPr = sec._sectPr
    cols = sectPr.xpath('./w:cols')
    if cols:
        print(f'  cols num={cols[0].get(qn("w:num"))}, space={cols[0].get(qn("w:space"))}')
    else:
        print('  no w:cols tag')

print('\n=== PARAGRAPHS ===')
for i, p in enumerate(doc.paragraphs[:45]):
    text_preview = p.text[:90].replace('\n', ' ')
    style_name = p.style.name if p.style else 'No Style'
    font_name = None
    font_size = None
    bold = None
    italic = None
    if p.runs:
        font_name = p.runs[0].font.name
        font_size = p.runs[0].font.size.pt if p.runs[0].font.size else None
        bold = p.runs[0].bold
        italic = p.runs[0].italic
    print(f'P{i:02d} [{style_name}] font={font_name} size={font_size} b={bold} it={italic}: {text_preview}')

print('\n=== TABLES ===')
print(f'Total tables: {len(doc.tables)}')
for idx, t in enumerate(doc.tables):
    print(f'Table {idx}: rows={len(t.rows)}, cols={len(t.columns)}')
    for r in t.rows[:3]:
        row_txt = [c.text.strip().replace('\n', ' ')[:30] for c in r.cells]
        print(f'  {row_txt}')
